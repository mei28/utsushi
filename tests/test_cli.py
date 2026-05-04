"""Smoke tests for the Typer CLI."""

from __future__ import annotations

from pathlib import Path

import pytest
from typer.testing import CliRunner

from utsushi import __version__
from utsushi.adapters import keynote
from utsushi.cli import app

runner = CliRunner()


def test_version_command_prints_version() -> None:
    result = runner.invoke(app, ["version"])
    assert result.exit_code == 0
    assert __version__ in result.stdout


def test_no_args_shows_help() -> None:
    result = runner.invoke(app, [])
    # no_args_is_help shows the help text; exit code can be 0 or 2 depending on Typer version
    assert "Usage" in result.stdout


def test_convert_command_invokes_pipeline(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    src = tmp_path / "deck.key"
    src.write_text("fake")

    def fake_export(s: Path, d: Path) -> None:
        d.write_bytes(b"PK\x03\x04")

    monkeypatch.setattr(keynote, "export_to_pptx", fake_export)
    result = runner.invoke(app, ["convert", str(src), "--to", "pptx"])
    assert result.exit_code == 0, result.stdout
    assert (tmp_path / "deck.pptx").exists()
    assert "deck.pptx" in result.stdout


def test_convert_with_explicit_out(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    src = tmp_path / "deck.key"
    src.write_text("fake")
    dst = tmp_path / "renamed.pptx"

    monkeypatch.setattr(keynote, "export_to_pptx", lambda s, d: d.write_bytes(b"PK\x03\x04"))
    result = runner.invoke(app, ["convert", str(src), "--to", "pptx", "--out", str(dst)])
    assert result.exit_code == 0, result.stdout
    assert dst.exists()


def test_convert_to_slides_invokes_drive_upload(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    from utsushi.adapters import drive

    src = tmp_path / "deck.pptx"
    src.write_bytes(b"PK\x03\x04")
    captured: dict[str, object] = {}

    def fake_upload(p: Path, folder_id: str | None = None) -> str:
        captured["path"] = p
        captured["folder_id"] = folder_id
        return "FILE_ID_X"

    monkeypatch.setattr(drive, "upload_as_slides", fake_upload)
    result = runner.invoke(app, ["convert", str(src), "--to", "slides"])
    assert result.exit_code == 0, result.stdout
    assert "FILE_ID_X" in result.stdout


def test_diff_command_identical_files_succeeds(tmp_path: Path) -> None:
    import shutil

    from pptx import Presentation as build_presentation

    a = tmp_path / "a.pptx"
    b = tmp_path / "b.pptx"
    prs = build_presentation()
    prs.slides.add_slide(prs.slide_layouts[0])
    prs.save(str(a))
    shutil.copyfile(a, b)

    result = runner.invoke(app, ["diff", str(a), str(b)])
    assert result.exit_code == 0, result.stdout
    assert "identical" in result.stdout.lower()


def test_diff_command_returns_nonzero_on_difference(tmp_path: Path) -> None:
    from pptx import Presentation as build_presentation

    a = tmp_path / "a.pptx"
    b = tmp_path / "b.pptx"
    prs = build_presentation()
    prs.slides.add_slide(prs.slide_layouts[0])
    prs.save(str(a))
    prs2 = build_presentation()
    prs2.slides.add_slide(prs2.slide_layouts[1])  # different layout
    prs2.save(str(b))

    result = runner.invoke(app, ["diff", str(a), str(b)])
    assert result.exit_code == 1


def test_normalize_command_invokes_apply_theme(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    from utsushi import normalize

    target = tmp_path / "deck.pptx"
    template = tmp_path / "tpl.pptx"
    target.write_bytes(b"PK\x03\x04")
    template.write_bytes(b"PK\x03\x04")
    captured: dict[str, object] = {}

    def fake_apply(t: Path, tpl: Path, out: Path | None = None, in_place: bool = False) -> Path:
        captured["target"] = t
        captured["template"] = tpl
        captured["out"] = out
        captured["in_place"] = in_place
        return out or t.with_name(f"{t.stem}.normalized{t.suffix}")

    monkeypatch.setattr(normalize, "apply_theme", fake_apply)
    result = runner.invoke(app, ["normalize", str(target), "--theme", str(template)])
    assert result.exit_code == 0, result.stdout
    assert captured["target"] == target
    assert captured["template"] == template
    assert captured["in_place"] is False
