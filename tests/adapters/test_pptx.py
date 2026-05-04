"""Unit tests for the PPTX adapter."""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation as build_presentation

from utsushi.adapters import pptx


@pytest.fixture
def tiny_pptx(tmp_path: Path) -> Path:
    """Create a minimal one-slide PPTX via python-pptx."""
    prs = build_presentation()
    prs.slides.add_slide(prs.slide_layouts[0])
    out = tmp_path / "tiny.pptx"
    prs.save(str(out))
    return out


def test_open_deck_returns_presentation_with_one_slide(tiny_pptx: Path) -> None:
    deck = pptx.open_deck(tiny_pptx)
    assert len(deck.slides) == 1


def test_open_deck_raises_for_missing_file(tmp_path: Path) -> None:
    with pytest.raises(FileNotFoundError):
        pptx.open_deck(tmp_path / "nope.pptx")


_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def test_get_theme_xml_returns_clrscheme(tiny_pptx: Path) -> None:
    deck = pptx.open_deck(tiny_pptx)
    theme = pptx.get_theme_xml(deck)
    assert theme.tag == f"{_NS}theme"
    assert theme.find(f".//{_NS}clrScheme") is not None
    assert theme.find(f".//{_NS}fontScheme") is not None


def test_set_theme_xml_then_save_round_trips(tiny_pptx: Path, tmp_path: Path) -> None:
    """Replacing the theme should survive a save/reload cycle."""
    from lxml import etree

    deck = pptx.open_deck(tiny_pptx)
    theme = pptx.get_theme_xml(deck)
    # Mutate the clrScheme name attribute to a sentinel.
    clrscheme = theme.find(f".//{_NS}clrScheme")
    assert clrscheme is not None
    clrscheme.set("name", "utsushi-test-sentinel")

    pptx.set_theme_xml(deck, theme)
    out = tmp_path / "out.pptx"
    pptx.save_deck(deck, out)

    reopened = pptx.open_deck(out)
    reopened_theme = pptx.get_theme_xml(reopened)
    reopened_clr = reopened_theme.find(f".//{_NS}clrScheme")
    assert reopened_clr is not None
    assert reopened_clr.get("name") == "utsushi-test-sentinel"

    # Sanity: confirm we serialized valid XML.
    etree.tostring(reopened_theme)
